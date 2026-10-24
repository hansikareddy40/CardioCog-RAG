"""Build an editable PowerPoint version of the final review deck (same content as final_review.tex).

Usage:  python presentation/make_pptx.py   ->  presentation/CardioCog_final_review.pptx
"""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
FIGS = [HERE / "figures", HERE.parent / "thesis" / "figures"]
NAVY, TEAL, TEALBG = RGBColor(0x1B, 0x2A, 0x4A), RGBColor(0x0E, 0x72, 0x71), RGBColor(0xE8, 0xF3, 0xF2)
MAROON, RED, GREY, BLACK = RGBColor(0x9C, 0x1D, 0x55), RGBColor(0xC6, 0x28, 0x28), RGBColor(0x6B, 0x72, 0x80), RGBColor(0, 0, 0)
W, H = 13.333, 7.5
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]
TOTAL = 27
state = {"n": 0}


def fig(name):
    for d in FIGS:
        if (d / name).exists():
            return d / name
    raise FileNotFoundError(name)


def text(slide, x, y, w, h, parts, size=16, color=BLACK, align=PP_ALIGN.LEFT, bold=False, anchor=MSO_ANCHOR.TOP):
    """parts: str, or list of paragraphs; a paragraph is a str or a list of (text, bold) runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.08)
    paras = parts if isinstance(parts, list) else [parts]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(4)
        runs = para if isinstance(para, list) else [(para, bold)]
        for t, b in runs:
            r = p.add_run()
            r.text = t
            r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), b, color, FONT
    return tb


def bullets(slide, x, y, w, h, items, size=16, numbered=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        mark = f"{i + 1}.  " if numbered else "•  "
        r = p.add_run()
        r.text = mark
        r.font.size, r.font.color.rgb, r.font.name = Pt(size), MAROON if numbered else NAVY, FONT
        runs = item if isinstance(item, list) else [(item, False)]
        for t, b in runs:
            r = p.add_run()
            r.text = t
            r.font.size, r.font.bold, r.font.name = Pt(size), b, FONT
    return tb


def rect(slide, x, y, w, h, fill, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    s.shadow.inherit = False
    return s


def block(slide, x, y, w, title, body, h=None, size=15):
    """Teal title strip over a light-teal body, as in the template."""
    lines = body if isinstance(body, list) else [body]
    est = sum(max(1, len(l if isinstance(l, str) else "".join(t for t, _ in l)) * size / 1100 / w * 1.35) for l in lines)
    h = h or 0.42 + 0.32 * est * (size / 15) + 0.15
    rect(slide, x, y, w, 0.42, TEAL)
    text(slide, x, y, w, 0.42, title, size=size + 1, color=RGBColor(255, 255, 255), anchor=MSO_ANCHOR.MIDDLE)
    rect(slide, x, y + 0.42, w, h - 0.42, TEALBG)
    text(slide, x, y + 0.46, w, h - 0.46, lines, size=size)
    return y + h


def finding(slide, x, y, w, title, body, size=15, h=1.2):
    text(slide, x, y, w, 0.4, title, size=size + 3, color=RED)
    text(slide, x, y + 0.42, w, h, body, size=size)


def image(slide, name, x, y, w=None, h=None):
    path = fig(name)
    iw, ih = Image.open(path).size
    if w and h:
        if iw / ih > w / h:
            h = w * ih / iw
        else:
            w = h * iw / ih
    elif w:
        h = w * ih / iw
    else:
        w = h * iw / ih
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    return x + w, y + h


def table(slide, x, y, w, rows, col_w=None, size=13, bold_cells=(), header=True, row_h=0.33):
    nr, nc = len(rows), len(rows[0])
    shp = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    t = shp.table
    tbl_pr = shp._element.graphic.graphicData.tbl
    tbl_pr.tblPr.set("bandRow", "0")
    tbl_pr.tblPr.set("firstRow", "0")
    if col_w:
        for j, cw in enumerate(col_w):
            t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor(0xF2, 0xF4, 0xF7) if (header and i == 0) else RGBColor(255, 255, 255)
            c.margin_top = c.margin_bottom = Inches(0.03)
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(val)
            r.font.size, r.font.name = Pt(size), FONT
            r.font.bold = (header and i == 0) or (i, j) in bold_cells
    return y + row_h * nr


def slide(title):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, W, 0.8, NAVY)
    text(s, 0.35, 0, W - 0.7, 0.8, title, size=26, color=RGBColor(255, 255, 255), anchor=MSO_ANCHOR.MIDDLE)
    state["n"] += 1
    text(s, W - 1.4, H - 0.45, 1.2, 0.35, f"{state['n']} / {TOTAL}", size=11, color=GREY, align=PP_ALIGN.RIGHT)
    return s


# ================================================================ title
s = prs.slides.add_slide(BLANK)
text(s, 1, 0.7, W - 2, 1.4, "CardioCog: An Explainable AI Framework for Cognitive Health Assessment and Early Prediction of Cognitive Decline Using Cardiovascular, Genetic, and Clinical Factors",
     size=26, color=MAROON, bold=True, align=PP_ALIGN.CENTER)
text(s, 1, 2.35, W - 2, 0.5, "FINAL REVIEW – B.TECH PROJECT", size=20, align=PP_ALIGN.CENTER)
text(s, 1, 2.95, W - 2, 0.45, "Team No.: D07", size=18, align=PP_ALIGN.CENTER)
names = [("Akula Akhil Vighnesh", "CB.SC.U4AIE23303"), ("Vada Gouri Hansika Reddy", "CB.SC.U4AIE23304"),
         ("Pilli Venkata Bhargav Reddy", "CB.SC.U4AIE23329"), ("Marella Venkata Mahidhar Reddy", "CB.SC.U4AIE23354")]
text(s, 3.4, 3.5, 3.9, 1.8, [n for n, _ in names], size=17)
text(s, 7.3, 3.5, 2.8, 1.8, [r for _, r in names], size=17)
text(s, 1, 5.25, W - 2, 0.45, "Supervisor: Dr. Manoj Bhat K, Professor", size=18, align=PP_ALIGN.CENTER)
text(s, 1, 5.95, W - 2, 0.8, "AMRITA VISHWA VIDYAPEETHAM – SCHOOL OF ARTIFICIAL INTELLIGENCE, COIMBATORE", size=17, align=PP_ALIGN.CENTER)

# ================================================================ 01 progress
s = slide("01 | Research Progress: From Proposal to Evidence")
y = block(s, 0.5, 1.05, 12.3, "Research progression",
          "Zeroth review: multimodal framework → first review: 3D amyloid-PET prototype (198 people, 17 events) → full NACC clinical cohort (19,470 people, 2,818 events) → two decline outcomes, five models, fusion network → PET image pipeline, prototype, 17 verification checks.", h=1.35)
table(s, 0.5, y + 0.25, 12.3, [
    ["Milestone", "Evidence and outcome"],
    ["First-review prototype", "3D CNN + cross-attention on amyloid PET: AUROC 0.786 in one run, 0.652 on re-run. Too few events to conclude."],
    ["Leakage-free cohort", "First visit, not demented, UDS v1–3; person-level split with 9 whole centres held out."],
    ["Dementia in 3 years", "XGBoost AUROC 0.938 (test), 0.946 (held-out centres); 0.84 within MCI."],
    ["Cognitive decline", "CDR-SB rise ≥ 1: AUROC 0.858 / 0.885; CDR-SB change regression helps MCI."],
    ["Independent checks", "17 of 17 checks pass; two future-information errors found and fixed."]],
    col_w=[3.0, 9.3], size=14, row_h=0.5)

# ================================================================ 02 problem
s = slide("02 | Motivation and Problem Statement")
y = block(s, 0.5, 1.05, 12.3, "Problem",
          "Clinics already record cognition, vascular history, genetics and sometimes PET for older adults without dementia, but rarely combine them into one explainable estimate of near-term decline. Which of these sources actually adds information, and can the estimate be trusted?", h=1.35)
text(s, 0.5, y + 0.15, 12.3, 0.9, ["For a person without dementia at visit t₀, estimate over the next three years:",
     "p̂_dem = P(dementia | x),     p̂_dec = P(ΔCDR-SB ≥ 1 | x),     Δ̂ = E(ΔCDR-SB | x)"], size=16)
bullets(s, 0.5, y + 1.1, 12.3, 3.5, [
    "How accurately can both outcomes be predicted, for normal and MCI groups separately?",
    "Does cardiovascular history add to demographics, genetics and current cognition?",
    "Do amyloid and tau PET add to the clinical model?",
    "Do deep models, including a modality-encoder fusion network, beat classical ones?",
    "Do earlier visits help?",
    "Can one model cope with whole categories of information missing?",
    "Can each estimate be explained and supported with published evidence?"], size=15, numbered=True)

# ================================================================ 03 literature
s = slide("03 | Literature Review and Research Gap")
y = table(s, 0.5, 1.05, 12.3, [
    ["Work", "Core method", "Relevance / limitation for CardioCog"],
    ["Livingston et al. (2020)", "Lancet Commission, 12 modifiable risk factors", "Vascular risk matters over the life course; mostly midlife exposure"],
    ["Alaka et al. (2025)", "CAIDE + heart rate, random forest, NACC", "Same database; AUC 0.80–0.91, no within-stage results"],
    ["Stern & Linial (2026)", "CatBoost on UK Biobank codes + APOE", "AUROC 0.773 without cognitive inputs"],
    ["Aschwanden et al. (2020)", "Random survival forests, 52 predictors", "Cardiovascular factors less important than expected"],
    ["Yi et al. (2023)", "XGBoost + SHAP for diagnosis", "CDR-SB most important; diagnosis, not prediction"],
    ["Gorishniy et al. (2021)", "FT-Transformer for tabular data", "No deep model consistently beats boosted trees"],
    ["Lewis et al. (2020)", "Retrieval-augmented generation", "Grounds generated text in a fixed library"]],
    col_w=[2.9, 4.3, 5.1], size=13, row_h=0.47)
block(s, 0.5, y + 0.3, 12.3, "Research gap",
      [[("Few studies measure, on one large leakage-controlled cohort, ", False), ("how much each source adds once current cognition is known", True),
        (", with results ", False), ("within", True), (" normal and MCI groups.", False)]], h=1.0)

# ================================================================ first review
s = slide("Why the design changed after the first review")
y = block(s, 0.5, 1.2, 5.9, "First-review prototype", "3D CNN on raw amyloid PET + 34 clinical features, spatial cross-attention, 1-year CDR-SB rise ≥ 0.5.", h=1.3)
table(s, 0.5, y + 0.3, 5.9, [["Run", "Mean AUROC (SD)"], ["Reported at review", "0.786 (0.134)"], ["Re-run, same notebook", "0.652 (0.153)"]], col_w=[3.4, 2.5], size=15, row_h=0.45)
finding(s, 7.0, 1.2, 5.8, "Diagnosis", [[("Only 198 people and ", False), ("17 events", True), (": a handful of participants decides the score. The pipeline also mixed PET time frames, had no registration, read one non-existent column and fitted SHAP on all data.", False)]], h=2.0)
finding(s, 7.0, 3.9, 5.8, "Decision", "Move to the full clinical cohort with three-year outcomes and thousands of events; treat PET images as a separate, validated pipeline.", h=1.5)

# ================================================================ 04 data
s = slide("04 | Data: NACC Freeze 74 and the Cohort")
y = table(s, 0.5, 1.1, 6.0, [["Step", "Participants"], ["All participants", "57,038"], ["First visit, in person, UDS v1–3", "55,461"],
                             ["   not demented at that visit", "38,107"], ["   outcome known at 3 years", "19,470"], ["      converted to dementia", "2,818 (14.5%)"]],
          col_w=[4.0, 2.0], size=15, row_h=0.45)
block(s, 0.5, y + 0.35, 6.0, "Three groups, very different risk", ["Normal 12,377 (1.6% convert)", "MCI 5,993 (41.5%)", "Impaired, not MCI 1,100 (12.3%)"], h=1.55)
image(s, "split.png", 6.9, 1.2, w=6.0)
bullets(s, 6.9, 3.0, 6.0, 3.0, [[("Person-level split; ", False), ("9 of 46 centres held out", True), (" entirely.", False)],
                                "Train 70% / validation 15% / test 15%, seed 2026.", "Every normaliser and imputer fitted on training data only."], size=16)

# ================================================================ outcomes
s = slide("Outcomes and censoring")
image(s, "outcome.png", 0.4, 1.3, w=6.9)
text(s, 7.6, 1.1, 5.3, 1.6, ["y = 1 if T ≤ 3", "y = 0 if T > 3 and C ≥ 3", "y unknown if T > 3 and C < 3",
                             ("T: years to dementia diagnosis; C: years to last visit.")], size=16)
y = block(s, 7.6, 2.95, 5.3, "Secondary outcomes", "CDR-SB rise ≥ 1 within 3 years (20,355 people, 5,035 events) and the 3-year change Δ in CDR-SB.", h=1.4)
finding(s, 7.6, y + 0.25, 5.3, "Robustness", "A Cox model that keeps censored participants gives the same conclusions.")

# ================================================================ 05 features
s = slide("05 | 48 Features in Five Categories")
y = table(s, 0.5, 1.1, 6.0, [["Category", "Features"], ["Demographics", "8"], ["Cardiovascular and medical", "25"], ["Cognitive tests (z-scores)", "9"],
                             ["Clinical stage (CDR-SB, diagnosis)", "3"], ["Genetics (APOE ε4, family history)", "3"], ["PET (second-stage model only)", "4"]],
          col_w=[4.6, 1.4], size=15, row_h=0.42)
block(s, 0.5, y + 0.3, 6.0, "Harmonising tests across form versions",
      ["ŝ = β₀ + β₁·age + β₂·female + β₃·education", "z = (s − ŝ) / σ_res",
       "Fitted on cognitively normal training participants."], h=1.6)
image(s, "zscore.png", 6.9, 1.1, w=6.0)
text(s, 6.9, 3.3, 6.0, 2.0, "MoCA example: expected 25.8 for a 72-year-old woman with 16 years of education, so a score of 21 gives z = −1.9. Healthy people not used for fitting score within 0.1 of zero in every UDS version.", size=16)

# ================================================================ 06 pipeline
s = slide("06 | Methodology: the CardioCog Pipeline")
_, y = image(s, "pipeline.png", 1.6, 0.95, w=10.1)
block(s, 0.5, y + 0.15, 12.3, "Rule for every step", "Only information available at the index visit; anything learned from data is fitted on the training set; results are reported on the test set and on centres never seen in training.", h=1.0)

# ================================================================ models
s = slide("Models compared on the same 48 features")
y = table(s, 0.5, 1.1, 6.2, [["Model", "Key settings"], ["Logistic regression", "L2 penalty, C from 5 values"], ["Random forest", "500 trees, √p features per split"],
                             ["XGBoost", "η = 0.03, depth 3–6, λ = 2, early stopping"], ["MLP", "128–64 units, dropout 0.2"], ["FT-Transformer", "64-d tokens, 3 layers, 8 heads"]],
          col_w=[2.3, 3.9], size=15, row_h=0.48)
text(s, 0.5, y + 0.2, 6.2, 1.2, "Three seeds averaged; XGBoost chosen as the main model (handles missing values, exact SHAP).", size=16)
y = block(s, 7.1, 1.1, 5.7, "XGBoost objective", ["f_M(x) = Σₘ η·gₘ(x)", "w*ⱼ = −Σ Gᵢ / (Σ Hᵢ + λ)",
                                                     "G, H: first and second derivatives of the loss"], h=1.75)
block(s, 7.1, y + 0.3, 5.7, "Loss and evaluation", ["L = −(1/n) Σ [y log p̂ + (1−y) log(1−p̂)]",
                                                     "AUROC, AUPRC, Brier, calibration slope, net benefit; 1,000 paired bootstrap resamples."], h=1.75)

# ================================================================ fusion network diagram
s = slide("Modality-encoder fusion network (proposed architecture)")
mods = ["Demographics", "Genetics", "Cardiovascular", "Cognitive tests", "Clinical stage"]
bw, gap, x0 = 2.2, 0.3, 0.55
centres = []
for i, m in enumerate(mods):
    x = x0 + i * (bw + gap)
    for yy, label, fill in [(1.15, m, TEALBG), (2.05, "Encoder E_g → token t_g", RGBColor(255, 255, 255))]:
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(yy), Inches(bw), Inches(0.6))
        b.fill.solid(); b.fill.fore_color.rgb = fill; b.line.color.rgb = NAVY; b.shadow.inherit = False
        tf = b.text_frame; tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = tf.paragraphs[0].add_run(); r.text = label; r.font.size = Pt(14); r.font.color.rgb = BLACK; r.font.name = FONT
    centres.append(x + bw / 2)


def boxed(x, y, w, h, label, fill, line):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.solid(); b.fill.fore_color.rgb = fill; b.line.color.rgb = line; b.shadow.inherit = False
    tf = b.text_frame; tf.word_wrap = True; tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tf.paragraphs[0].add_run(); r.text = label; r.font.size = Pt(14); r.font.color.rgb = BLACK; r.font.name = FONT
    return b


def arrow(x1, y1, x2, y2, color=NAVY):
    c = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(1.5)
    c.line._get_or_add_ln().append(c.line._get_or_add_ln().makeelement(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", {"type": "triangle"}))


att = boxed(3.65, 3.25, 6.0, 0.6, "4-head attention with learned query q → fused z", TEALBG, NAVY)
for c in centres:
    arrow(c, 1.75, c, 2.05)
    arrow(c, 2.65, 6.65, 3.25)
boxed(2.4, 4.4, 3.6, 0.6, "Risk head: p̂ = σ(r(z))", RGBColor(0xF5, 0xE8, 0xEE), MAROON)
boxed(7.3, 4.4, 3.6, 0.6, "Progression head: Δ̂ = q(z)", RGBColor(0xF5, 0xE8, 0xEE), MAROON)
boxed(11.2, 4.4, 1.85, 0.8, "Tier 2: PET by late fusion", RGBColor(255, 255, 255), TEAL)
arrow(6.65, 3.85, 4.2, 4.4)
arrow(6.65, 3.85, 9.1, 4.4)
text(s, 0.5, 5.4, 12.3, 1.5, ["A missing modality gets a learned token m_g; each token gets a type embedding τ_g; whole modalities are hidden with probability 0.35 in training.",
                             "L = L_BCE + 0.5 · Huber(Δ − Δ̂),     α_g = softmax_g((W_Q q)ᵀ W_K t_g / √8)"], size=16, align=PP_ALIGN.CENTER)

# ================================================================ 07 results
s = slide("07 | Results: Dementia Within Three Years")
y = table(s, 0.5, 1.05, 12.3, [
    ["Model", "Test AUROC (95% CI)", "AUPRC", "Brier", "Held-out AUROC", "Held-out Brier"],
    ["Logistic regression", "0.936 (0.924–0.949)", "0.721", "0.066", "0.945", "0.060"],
    ["Random forest", "0.932 (0.919–0.944)", "0.704", "0.070", "0.940", "0.064"],
    ["XGBoost", "0.938 (0.925–0.948)", "0.713", "0.066", "0.946", "0.059"],
    ["MLP", "0.938 (0.926–0.949)", "0.721", "0.066", "0.946", "0.059"],
    ["FT-Transformer", "0.939 (0.927–0.949)", "0.720", "0.066", "0.945", "0.061"]],
    col_w=[2.8, 2.9, 1.4, 1.4, 2.0, 1.8], size=15, bold_cells={(3, 1), (3, 4)}, row_h=0.45)
text(s, 0.5, y + 0.05, 12.3, 0.4, "Test set: 2,286 participants, 332 events. Held-out centres: 4,064 participants, 545 events.", size=13, color=GREY)
y = block(s, 0.5, y + 0.5, 12.3, "What the result does establish", "All five models discriminate well and are well calibrated (constant-prediction Brier 0.124), at new centres as well as on the test set.", h=1.0)
block(s, 0.5, y + 0.2, 12.3, "What it does not establish", [[("Intervals overlap almost completely: deep models give ", False), ("no gain", True), (" over logistic regression or boosted trees.", False)]], h=1.0)

# ================================================================ by group
s = slide("Is 0.94 too good? Results within diagnostic groups")
y = table(s, 0.5, 1.2, 6.2, [["Group (test / held-out events)", "Test", "Held-out"], ["All not demented", "0.937", "0.946"],
                             ["Cognitively normal (24 / 37)", "0.884", "0.876"], ["MCI (289 / 496)", "0.835", "0.846"]],
          col_w=[3.8, 1.2, 1.2], size=16, row_h=0.5)
block(s, 0.5, y + 0.4, 6.2, "Interpretation", [[("Diagnosis alone gives AUROC 0.84, because MCI converts 25 times as often as normal cognition. Within MCI the model still ranks risk at ", False), ("0.84", True), (".", False)]], h=1.5)
finding(s, 7.2, 1.2, 5.6, "Critical check", [[("Removing all current-cognition information drops AUROC from 0.937 to ", False), ("0.746", True), (".", False)]], h=1.1)
finding(s, 7.2, 3.0, 5.6, "Meaning", "The model estimates short-term progression given the current clinical stage. It is not leakage: inputs are from the index visit, the label from later visits.", h=1.8)

# ================================================================ ablation
s = slide("What drives the prediction? Category ablation")
_, y = image(s, "06_feature_groups.png", 0.9, 0.95, w=11.5)
table(s, 2.4, y + 0.1, 8.5, [["Category", "Alone", "Gain added last (test)", "Gain (held-out)"],
                             ["Cognitive tests", "0.896", "+0.022", "+0.019"], ["Clinical stage", "0.897", "+0.015", "+0.021"],
                             ["Demographics", "0.671", "+0.010", "+0.009"], ["Genetics (APOE)", "0.645", "+0.003", "+0.001"],
                             ["Cardiovascular / medical", "0.626", "+0.001", "+0.001"]], col_w=[3.0, 1.3, 2.3, 1.9], size=14, row_h=0.36)

# ================================================================ cardio
s = slide("The cardiovascular question (RQ2)")
image(s, "09_hazard_ratios.png", 0.4, 1.3, w=6.5)
bullets(s, 7.2, 1.1, 5.7, 2.5, [[("Alone: AUROC ", False), ("0.626", True), ("; before cognition: +0.024 (test), +0.004 (held-out).", False)],
                                "After cognition: +0.001 for dementia, +0.0004 for CDR-SB decline.",
                                "Cox model: hypertension HR 0.995, diabetes and stroke ≈ 1."], size=16)
block(s, 7.2, 3.9, 5.7, "Interpretation", "Limited incremental value in this cohort, not evidence that vascular health is unimportant: median age 71, midlife exposure unobserved, and past damage already shows in the cognitive scores.", h=2.0)

# ================================================================ SHAP
s = slide("Explainability: Shapley values")
image(s, "07_shap_features.png", 0.3, 1.1, w=7.6)
text(s, 8.2, 1.0, 4.7, 0.5, "f(x) = φ₀ + Σⱼ φⱼ", size=18, align=PP_ALIGN.CENTER)
y = table(s, 8.4, 1.6, 4.3, [["Category", "SHAP share"], ["Clinical stage", "40.8%"], ["Cognitive tests", "35.0%"], ["Demographics", "15.5%"],
                             ["Cardiovascular", "5.1%"], ["Genetics", "3.5%"]], col_w=[2.6, 1.7], size=15, row_h=0.42)
text(s, 8.2, y + 0.2, 4.7, 1.8, "Stable over 8 refits (Spearman 0.89–0.97). Low BMI raises risk: weight loss before diagnosis, not causal.", size=15)

# ================================================================ decline
s = slide("Cognitive decline as the outcome")
_, y = image(s, "decline_vs_dementia.png", 1.9, 0.95, w=9.5)
block(s, 0.5, y + 0.1, 6.0, "CDR-SB rise ≥ 1 (20,355 / 5,035 events)", "AUROC 0.858 test, 0.885 held-out; normal 0.76–0.83 with 129 and 215 events, so more reliable than the dementia estimate for normal cognition.", h=1.65, size=14)
block(s, 6.8, y + 0.1, 6.0, "Predicted 3-year CDR-SB change", "MAE 0.69 → 0.60 (test), 0.82 → 0.72 (held-out); MCI 1.90 → 1.59; normal 0.17 = 0.17, no gain.", h=1.65, size=14)

# ================================================================ fusion results
s = slide("Fusion network: works, but no gain and uninformative attention")
image(s, "fusion_attention.png", 0.3, 1.3, w=7.4)
y = table(s, 8.0, 1.1, 4.8, [["AUROC", "Fusion net", "XGBoost"], ["Test", "0.934", "0.937"], ["Held-out", "0.943", "0.946"], ["MCI (test)", "0.831", "0.835"]],
          col_w=[1.8, 1.5, 1.5], size=15, row_h=0.42)
text(s, 7.9, y + 0.15, 5.0, 1.0, "Progression head MAE 0.72 → 0.65. Without cognition: AUROC 0.72, Brier 0.114.", size=15)
finding(s, 7.9, y + 1.25, 5.0, "Finding", "Attention is 18–22% for every modality (uniform = 20%), so it does not explain predictions. SHAP and ablation are used instead.", h=1.6)

# ================================================================ history and PET
s = slide("Earlier visits and PET measures (RQ3, RQ5)")
_, y1 = image(s, "06_longitudinal.png", 0.4, 1.1, w=6.1)
block(s, 0.4, y1 + 0.2, 6.1, "Visit history", "Landmark visits with ≥ 2 earlier visits: trajectory features +0.003, GRU +0.000 to +0.005 AUROC.", h=1.3)
_, y2 = image(s, "08_conversion_by_amyloid.png", 6.8, 1.1, w=6.1)
block(s, 6.8, y2 + 0.2, 6.1, "Amyloid and tau PET", "MCI: 34% of amyloid-positive vs 7% of negative convert within 2 years. Late fusion: 0.964 vs 0.944 on 243 people (33 events); gain interval −0.009 to +0.066 includes zero.", h=1.9)

# ================================================================ 08 PET images
s = slide("08 | PET Image Pipeline: From DICOM to SUVR")
_, y = image(s, "11_group_average_images.png", 0.4, 1.0, w=6.4, h=5.6)
text(s, 0.4, y + 0.02, 6, 0.35, "Group averages only (318 amyloid, 232 tau, 49 FDG scans).", size=11, color=GREY)
bullets(s, 7.1, 1.0, 5.8, 3.4, [[("Convert: ", True), ("599 dynamic DICOM series → static volume, V = Σ dₖVₖ / Σ dₖ over the standard window.", False)],
                                [("Register: ", True), ("rigid to MNI152, then 12-parameter affine to a study PET template (Mattes mutual information).", False)],
                                [("Quantify: ", True), ("Centiloid masks, SUVR = mean cortex / mean cerebellum.", False)]], size=15, numbered=True)
block(s, 7.1, 4.5, 5.8, "Validation against the NACC PET core", [[("Pearson r = ", False), ("0.977", True), (" over 318 scans (0.96–0.99 per tracer); mean difference −0.043.", False)]], h=1.3)

s = slide("3D CNN for amyloid status and Grad-CAM")
image(s, "11_gradcam_mean.png", 0.4, 1.0, w=6.5, h=5.7)
y = table(s, 7.2, 1.0, 5.7, [["Method (292 scans, 84 positive)", "AUROC"], ["Cortical SUVR (one number)", "0.982"], ["Regional logistic model", "0.971"], ["3D CNN (494,917 weights)", "0.960"]],
          col_w=[4.3, 1.4], size=15, bold_cells={(1, 1)}, row_h=0.42)
y = block(s, 7.2, y + 0.3, 5.7, "Interpretation", "The pipeline supports a working image model, but the label is itself derived from cortical SUVR, so one number is hard to beat.", h=1.4)
finding(s, 7.2, y + 0.2, 5.7, "Limitation", "Grad-CAM was diffuse (1.7% of top voxels in the cortical region vs 3.8% by chance), and only 5 imaged people converted, so images were not used to predict dementia.", h=1.8)

# ================================================================ 09 validation
s = slide("09 | Validation: Calibration, Usefulness, Missing Data")
_, y = image(s, "09_decision_curve.png", 0.4, 1.0, w=6.2)
bullets(s, 0.4, y + 0.1, 6.2, 3.0, ["Test: predicted 14.5% = observed 14.5%, slope 1.01.", "Held-out: 15.0% vs 13.4%, slope 1.10.",
                                    "At 20% risk: sensitivity 83%, specificity 88%, NPV 97%.", "Net benefit above treat-all and treat-none from 2% to 60%."], size=15)
y = table(s, 6.9, 1.0, 6.0, [["Brier score", "Standard", "Dropout", "Specialist"], ["Everything", "0.066", "0.067", "0.066"],
                             ["No clinical stage", "0.185", "0.076", "0.074"], ["No cognition at all", "0.441", "0.115", "0.111"], ["Demographics only", "0.444", "0.120", "0.118"]],
          col_w=[2.4, 1.2, 1.2, 1.2], size=14, row_h=0.42)
block(s, 6.9, y + 0.3, 6.0, "Modality dropout (p = 0.35)", "One model stays close to a specialist retrained for each missing pattern; a standard model becomes worse than the base rate.", h=1.4)

# ================================================================ 10 prototype
s = slide("10 | Research Prototype for Clinicians")
image(s, "app_estimate.png", 0.6, 0.95, h=6.1)
text(s, 6.3, 1.1, 6.6, 0.9, [[("Example: ", True), ("72-year-old woman, 16 years of education, MCI, CDR-SB 2, MoCA 21.", False)]], size=16)
bullets(s, 6.3, 2.1, 6.6, 4.5, ["FastAPI + React; form generated from the 48 trained features.", "Risk shown next to the observed rate of similar people (39% vs 40%), with SHAP by category.",
                                "Plausibility checks flag contradictory entries.", "Plain-language report with quoted evidence.",
                                "27 of 28 behaviour checks pass; MCI under 65 over-predicted (35% vs 27%)."], size=16)

s = slide("Plain-language report for the same patient")
image(s, "app_report.png", 0.6, 0.95, h=6.1)
bullets(s, 5.5, 1.1, 7.4, 3.4, [[("Summary: ", True), ("risk next to the observed rate of similar people.", False)], [("Main reasons: ", True), ("largest SHAP contributions in plain words.", False)],
                                [("Reliability: ", True), ("missing categories and plausibility warnings.", False)], [("Published evidence: ", True), ("quoted passages with their source.", False)],
                                [("Limits: ", True), ("fixed statement that it is not a diagnosis.", False)]], size=16, numbered=True)
block(s, 5.5, 4.7, 7.4, "Design choice", "Templates are filled with the person's own numbers and evidence is quoted, so the report cannot state anything the model or the literature did not provide.", h=1.4)

s = slide("Evidence assistant (retrieval-augmented generation)")
_, y = image(s, "app_ask.png", 1.4, 0.95, w=10.5)
text(s, 0.5, y + 0.2, 6.2, 2.2, ["s_RRF(d) = 1/(60 + rank_sem(d)) + 1/(60 + rank_BM25(d))",
                                 "MiniLM embeddings + BM25 (k₁ = 1.5, b = 0.75), merged by rank fusion. Below similarity 0.35 the assistant says the library has no evidence."], size=15)
table(s, 7.0, y + 0.2, 5.8, [["Test (11 papers, 447 passages)", "Result"], ["Correct source first: hybrid", "32 / 32"], ["   embeddings only / BM25 only", "28 / 31"],
                             ["Personal-advice questions flagged", "8 / 8"], ["Out-of-scope questions declined", "6 / 6"]], col_w=[4.4, 1.4], size=14, row_h=0.38)

# ================================================================ 11 verification
s = slide("11 | Independent Verification")
y = table(s, 0.5, 1.05, 12.3, [["Check", "Result"], ["No person or held-out centre in more than one partition", "0 shared"],
                               ["Features unchanged when all later visits are deleted", "Identical for 1,500 people × 67 features"],
                               ["No PET value newer than the visit that uses it", "0 of 1,481 amyloid, 713 tau"],
                               ["Labels recomputed independently", "0 disagreements in 11,403 visits"], ["Shuffled labels", "AUROC 0.508 (chance)"],
                               ["Retrain with a new seed", "0.937 / 0.946 (reported 0.938 / 0.946)"]], col_w=[7.0, 5.3], size=15, row_h=0.45)
block(s, 0.5, y + 0.3, 12.3, "Errors the checks found and we fixed",
      [[("Education and race filled from a later visit (0.1% of people); PET scans matched to visits weeks before the scan. After fixing, every model was retrained: results moved only in the third decimal. ", False), ("17 of 17 checks pass.", True)]], h=1.35)

# ================================================================ synthesis
s = slide("Synthesis: what each research question found")
y = table(s, 0.5, 1.05, 12.3, [["Research question", "Answer and evidence"],
                               ["RQ1 Prediction accuracy", "Dementia 0.94 (0.84 in MCI); decline 0.86 (0.75 in MCI); change regression helps MCI only"],
                               ["RQ2 Cardiovascular history", "+0.001 AUROC once cognition is known; Cox agrees"],
                               ["RQ3 Amyloid and tau PET", "Strong association, no demonstrable gain with current follow-up"],
                               ["RQ4 Deep models / fusion net", "No gain over logistic regression or XGBoost; attention uniform"],
                               ["RQ5 Earlier visits", "+0.000 to +0.005 AUROC"], ["RQ6 Missing categories", "Modality dropout keeps calibration close to specialist models"],
                               ["RQ7 Explanation and evidence", "SHAP, plausibility checks, report with quoted published evidence"]], col_w=[3.6, 8.7], size=15, row_h=0.47)
block(s, 0.5, y + 0.3, 12.3, "Final takeaway", "How a person is functioning today carries most of the three-year signal; decades of vascular and genetic history leave their trace in that measurement rather than adding to it.", h=1.05)

# ================================================================ limitations
s = slide("Limitations, Future Work and Final Takeaway")
bullets(s, 0.6, 1.0, 12.2, 4.6, ["Held-out centres are site-level validation within NACC; external and temporal validation on another cohort is the next step.",
                                 "NACC is a research cohort, mostly White and highly educated; about half of eligible people lacked a 3-year outcome (Cox check agrees).",
                                 "Only 24 and 37 cognitively normal converters in test and held-out sets; a dedicated 5- to 10-year model for normal cognition is needed.",
                                 "Richer vascular exposure (repeated blood pressure, treatment, small-vessel imaging) may capture what yes/no history misses.",
                                 "PET registration without MRI, tau not quantified; more follow-up will allow a powered PET test.",
                                 "Evidence library of 11 papers and team-written questions; a curated library and stronger local model would help."], size=16)
block(s, 0.6, 5.55, 12.2, "Contribution", "A leakage-free, verified pipeline that measures what each source of information adds, reports results within diagnostic groups, and delivers explainable, evidence-backed estimates.", h=1.05)

# ================================================================ thanks
s = prs.slides.add_slide(BLANK)
text(s, 1, 2.2, W - 2, 1.2, "Thank you", size=54, color=MAROON, bold=True, align=PP_ALIGN.CENTER)
text(s, 1, 3.5, W - 2, 0.6, "Questions are welcome.", size=22, align=PP_ALIGN.CENTER)
text(s, 1, 4.8, W - 2, 1.0, ["Data: National Alzheimer's Coordinating Center (NACC), Freeze 74, used under the NACC data use agreement.", "All results are aggregates."],
     size=14, color=GREY, align=PP_ALIGN.CENTER)

assert state["n"] == TOTAL, state["n"]
out = HERE / "CardioCog_final_review.pptx"
prs.save(out)
print("wrote", out, len(prs.slides), "slides")
