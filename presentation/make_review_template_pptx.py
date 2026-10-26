"""Build the final review deck in a plain version of the team's zeroth-review format.

The title and closing slides reuse the Amrita title slide from
presentation/template/D07_zeroth_review.pptx (4:3, Calibri). Every other slide uses one
frame: black title top left, a thin grey line, dark text, light-grey tables. The content
is the same as final_review.tex and the thesis.

Usage:  python presentation/make_review_template_pptx.py
        ->  presentation/CardioCog_final_review_D07.pptx
"""
import copy
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "template" / "D07_zeroth_review.pptx"
OUT = HERE / "CardioCog_final_review_D07.pptx"
FIGS = [HERE / "figures", HERE.parent / "thesis" / "figures"]
INK = RGBColor(0x1F, 0x1F, 0x1F)
GREY = RGBColor(0x59, 0x59, 0x59)
RULE = RGBColor(0xBF, 0xBF, 0xBF)
NO_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # built-in "No Style, No Grid"; rules are drawn per cell
X0, W0 = 0.5, 9.0  # content margins on the 10 x 7.5 in slide
TOP = 1.4  # content starts below the title and line

prs = Presentation(str(TEMPLATE))
REF = list(prs.slides)
N_REF = len(REF)
BLANK = prs.slide_layouts[6]
state = {"n": 0}


# ---------------------------------------------------------------- helpers
def fig(name):
    for d in FIGS:
        if (d / name).exists():
            return d / name
    raise FileNotFoundError(name)


def clone(i):
    """New slide copying reference slide i (1-based), pictures included."""
    src = REF[i - 1]
    dst = prs.slides.add_slide(src.slide_layout)
    for ph in list(dst.placeholders):
        ph._element.getparent().remove(ph._element)
    for sh in src.shapes:
        el = copy.deepcopy(sh._element)
        for blip in el.iter(qn("a:blip")):
            part = src.part.related_part(blip.get(qn("r:embed")))
            blip.set(qn("r:embed"), dst.part.relate_to(part, RT.IMAGE))
        dst.shapes._spTree.append(el)
    return dst


def shape(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(name)


def place(sh, x, y, w, h):
    """Set all four: a placeholder that inherits its frame gets zeros for any side left unset."""
    sh.left, sh.top, sh.width, sh.height = Inches(x), Inches(y), Inches(w), Inches(h)


def runs(p, text, size=None, color=INK, bold=None, italic=None):
    """Add runs to paragraph p; **...** marks bold."""
    for k, seg in enumerate(text.split("**")):
        if not seg:
            continue
        r = p.add_run()
        r.text = seg
        f = r.font
        f.name = "Calibri"
        if size:
            f.size = Pt(size)
        if color is not None:
            f.color.rgb = color
        f.bold = bool(k % 2 == 1 or bold)
        if italic:
            f.italic = True


def set_bullet(p, on, level=0):
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    if on:
        pPr.set("marL", str(Inches(0.25 + 0.3 * level)))
        pPr.set("indent", str(-Inches(0.22)))
        etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
        etree.SubElement(pPr, qn("a:buChar")).set("char", "–" if level else "•")
    else:
        pPr.set("marL", "0")
        pPr.set("indent", "0")
        etree.SubElement(pPr, qn("a:buNone"))


def fill(tf, paras, size=16, color=INK, space=6, align=None, bullets=True):
    """paras: list of str or dict(text, h=heading, lvl, size, color, bold, italic, gap, plain).
    A plain str is a bullet (if bullets); h=True is an unbulleted bold heading."""
    txBody = tf._txBody
    bp = txBody.find(qn("a:bodyPr"))
    for e in list(bp):
        if e.tag in (qn("a:normAutofit"), qn("a:spAutoFit")):
            bp.remove(e)
    for p in txBody.findall(qn("a:p")):
        txBody.remove(p)
    tf.word_wrap = True
    for item in paras:
        d = {"text": item} if isinstance(item, str) else dict(item)
        p = tf.add_paragraph()
        heading = d.get("h", False)
        set_bullet(p, bullets and not heading and not d.get("plain"), d.get("lvl", 0))
        p.space_after = Pt(d.get("gap", space))
        if heading:
            p.space_before = Pt(d.get("before", 6))
        if align is not None:
            p.alignment = align
        runs(p, d["text"], size=d.get("size", size), color=d.get("color", color),
             bold=heading or d.get("bold"), italic=d.get("italic"))
    return tf


def box(slide, x, y, w, h, paras, size=16, bullets=True, align=None, **kw):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(0.04))
    fill(tf, paras, size=size, bullets=bullets, align=align, **kw)
    return tb


def note(slide, x, y, w, h, heading, text, size=14):
    """Bold heading, then a plain paragraph."""
    return box(slide, x, y, w, h, [{"text": heading, "h": True, "gap": 2, "before": 0},
                                   {"text": text, "plain": True}], size=size, bullets=False)


def caption(slide, x, y, w, text):
    box(slide, x, y, w, 0.35, [{"text": text, "plain": True, "color": GREY}], size=11, bullets=False)


def pic(slide, name, x, y, w, h):
    """Picture fitted inside the box (x, y, w, h), centred."""
    path = fig(name)
    iw, ih = Image.open(path).size
    s = min(w / iw, h / ih)
    pw, ph = iw * s, ih * s
    return slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
                                    Inches(pw), Inches(ph))


def rules(cell, header):
    """Horizontal grey rules only; the header row gets a light grey fill."""
    tcPr = cell._tc.get_or_add_tcPr()
    for side, show in (("a:lnL", False), ("a:lnR", False), ("a:lnT", header), ("a:lnB", True)):
        ln = etree.SubElement(tcPr, qn(side))
        if show:
            ln.set("w", str(Pt(1.25 if header else 0.75)))
            fill = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(fill, qn("a:srgbClr")).set("val", "7F7F7F" if header else "D0D0D0")
        else:
            ln.set("w", "0")
            etree.SubElement(ln, qn("a:noFill"))
    if header:
        fill = etree.SubElement(tcPr, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr")).set("val", "F2F2F2")


def table(slide, x, y, w, cols, rows, size=13, row_h=0.36, bold_col0=False, bold_cells=()):
    """rows[0] is the header; cols are relative widths."""
    gf = slide.shapes.add_table(len(rows), len(cols), Inches(x), Inches(y), Inches(w),
                                Inches(row_h * len(rows)))
    gf.table._tbl.tblPr.find(qn("a:tableStyleId")).text = NO_STYLE
    t = gf.table
    tot = sum(cols)
    for j, c in enumerate(cols):
        t.columns[j].width = Inches(w * c / tot)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text_frame.word_wrap = True
            rules(cell, i == 0)
            runs(cell.text_frame.paragraphs[0], val, size=size,
                 bold=(i == 0) or (bold_col0 and j == 0) or (i, j) in bold_cells)
    return gf


def frame(t):
    """The one content frame: title top left, thin grey line, slide number bottom right."""
    s = prs.slides.add_slide(BLANK)
    state["n"] += 1
    tb = s.shapes.add_textbox(Inches(X0), Inches(0.35), Inches(W0), Inches(0.75))
    tf = tb.text_frame
    tf.vertical_anchor = MSO_ANCHOR.BOTTOM
    tf.margin_left = tf.margin_right = Inches(0)
    fill(tf, [{"text": t, "plain": True}], size=30, bullets=False)
    line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(X0), Inches(1.15), Inches(X0 + W0), Inches(1.15))
    line.line.color.rgb = RULE
    line.line.width = Pt(1)
    box(s, 8.9, 7.05, 0.6, 0.3, [{"text": str(state["n"] + 1), "plain": True, "color": GREY}],
        size=10, bullets=False, align=PP_ALIGN.RIGHT)
    return s


# ================================================================ title
s = clone(1)
retitle = shape(s, "Title 1")
place(retitle, 0.6, 1.4, 8.8, 1.8)
fill(retitle.text_frame, [{"text": "CardioCog: An Explainable AI Framework for Cognitive Health Assessment and "
                                   "Early Prediction of Cognitive Decline Using Cardiovascular, Genetic, and "
                                   "Clinical Factors", "plain": True}], size=22, bullets=False, align=PP_ALIGN.CENTER)
sub = shape(s, "Subtitle 2")
place(sub, 0.5, 3.4, 9.0, 3.8)
fill(sub.text_frame, [
    {"text": "Final Review – B.Tech Project", "bold": True, "size": 18, "color": INK, "gap": 12},
    {"text": "Akula Akhil Vighnesh – CB.SC.U4AIE23303", "gap": 0},
    {"text": "Vada Gouri Hansika Reddy – CB.SC.U4AIE23304", "gap": 0},
    {"text": "Pilli Venkata Bhargav Reddy – CB.SC.U4AIE23329", "gap": 0},
    {"text": "Marella Venkata Mahidhar Reddy – CB.SC.U4AIE23354", "gap": 12},
    {"text": "Team D07", "gap": 4},
    {"text": "Guide: Dr. Manoj Bhat K, Professor  |  Co-guide: Amrutha Veluppal, Professor", "gap": 0},
], size=14, color=GREY, bullets=False, align=PP_ALIGN.CENTER)

# ================================================================ introduction
s = frame("Introduction & Motivation")
box(s, X0, TOP, W0, 5.6, [
    {"text": "Domain: AI / Healthcare Analytics / XAI / Biomedical Data Science", "plain": True, "color": GREY},
    {"text": "Introduction", "h": True},
    "Dementia and Alzheimer's disease are increasing worldwide, affecting memory, thinking and quality of life.",
    "Hypertension, diabetes, obesity and stroke are associated with a higher risk of cognitive decline.",
    "Clinics already record cognition, vascular history, genetics and sometimes PET, but rarely combine them into one explainable estimate of near-term risk.",
    {"text": "Motivation", "h": True},
    "Decline is hard to reverse once noticeable, so the years before diagnosis are the window for intervention.",
    "Early identification allows low-cost steps and gives families time to plan.",
    "In NACC, **14.5%** of people without dementia converted within three years; among those with MCI, **41.5%**.",
], size=17)

# ================================================================ problem statement
s = frame("Problem Statement")
box(s, X0, TOP, W0, 1.9, [
    {"text": "Clinical, cardiovascular, genetic and PET information is rarely combined into one explainable "
             "estimate of near-term decline, and it is unclear which source actually adds information.", "plain": True},
    {"text": "For a person without dementia at a visit, CardioCog estimates over the next three years:", "plain": True, "gap": 2},
    {"text": "the risk of **dementia**", "gap": 0},
    {"text": "the risk of a **CDR-SB rise ≥ 1**", "gap": 0},
    "the expected **change in CDR-SB**",
], size=16)
table(s, X0, 4.0, W0, [1, 1], [
    ["In scope", "Out of scope"],
    ["3-year dementia and CDR-SB decline risk", "Clinical diagnosis"],
    ["Demographic, cardiovascular, genetic, cognitive, clinical-stage and PET data", "Treatment recommendation"],
    ["SHAP explanations and quoted published evidence (RAG)", "Post-diagnosis prognosis"],
    ["Validation at 9 held-out NACC centres", "Validation on a non-NACC cohort"],
], size=14, row_h=0.6)

# ================================================================ research questions
s = frame("Research Questions")
table(s, X0, TOP + 0.1, W0, [0.7, 8.3], [
    ["RQ", "Question"],
    ["1", "How accurately can dementia and CDR-SB decline be predicted, within normal and MCI groups separately?"],
    ["2", "Does cardiovascular history add to demographics, genetics and current cognition?"],
    ["3", "Do amyloid and tau PET add to the clinical model?"],
    ["4", "Do deep models, including a modality-encoder fusion network, beat classical ones?"],
    ["5", "Do earlier visits help?"],
    ["6", "Can one model cope with whole categories of missing information?"],
    ["7", "Can each estimate be explained and supported with published evidence?"],
], size=15, row_h=0.62)

# ================================================================ literature
s = frame("Literature Review")
table(s, X0, TOP + 0.1, W0, [2.2, 3.1, 3.7], [
    ["Work", "Core method", "Relevance / limitation"],
    ["Livingston et al. (2020)", "Lancet Commission, 12 modifiable risk factors", "Vascular risk matters over the life course; mostly midlife exposure"],
    ["Alaka et al. (2025)", "CAIDE + resting heart rate, random forest, NACC", "Same database; AUC 0.80–0.91, no within-stage results"],
    ["Stern & Linial (2026)", "CatBoost on UK Biobank codes + APOE", "AUROC 0.773 without cognitive inputs"],
    ["Aschwanden et al. (2020)", "Random survival forests, 52 predictors", "Cardiovascular factors less important than expected"],
    ["Yi et al. (2023)", "XGBoost + SHAP for diagnosis", "CDR-SB most important; diagnosis, not prediction"],
    ["Gorishniy et al. (2021)", "FT-Transformer for tabular data", "No deep model consistently beats boosted trees"],
    ["Lewis et al. (2020)", "Retrieval-augmented generation", "Grounds generated text in a fixed library"],
], size=13, row_h=0.66)

# ================================================================ research gap
s = frame("Research Gap and Our Approach")
box(s, X0, TOP, W0, 2.8, [
    {"text": "Research gap", "h": True, "before": 0},
    "Few studies measure what cardiovascular factors add **once current cognition is known**.",
    "Results are rarely reported **within** normal and MCI groups, where the clinical question lies.",
    "Many models lack held-out sites and checks for information leakage.",
], size=17)
box(s, X0, 3.6, W0, 2.8, [
    {"text": "Our approach", "h": True, "before": 0},
    "One leakage-controlled NACC cohort: 19,470 people, nine whole centres held out.",
    "Category-by-category ablation, with results within diagnostic groups.",
    "17 independent verification checks on the whole pipeline.",
], size=17)

# ================================================================ objectives
s = frame("Objectives")
box(s, X0, TOP, W0, 5.6, [
    "Build a leakage-free three-year cohort from NACC with person-level splits and nine whole centres held out.",
    "Predict dementia and CDR-SB decline, and estimate the three-year CDR-SB change, from 48 features in five categories.",
    "Measure what cardiovascular history, genetics, PET and earlier visits add once current cognition is known.",
    "Compare classical models, deep tabular models and a modality-encoder fusion network that copes with missing categories.",
    "Explain each estimate with SHAP and quoted published evidence in a research prototype, and verify the pipeline independently.",
], size=17, space=14)

# ================================================================ first review
s = frame("Changes After the First Review")
note(s, X0, TOP, 4.3, 2.6, "First-review prototype",
     "3D CNN on raw amyloid PET + 34 clinical features; outcome: 1-year CDR-SB rise ≥ 0.5. "
     "Mean AUROC 0.786 at the review, 0.652 when the same notebook was re-run.", size=15)
note(s, 5.2, TOP, 4.3, 2.6, "Diagnosis",
     "Only 198 people and 17 events, so a handful of participants decides the score. PET time frames were mixed, "
     "with no registration, and SHAP was fitted on all data.", size=15)
note(s, X0, 3.7, W0, 2.0, "Decision",
     "Move to the full clinical cohort with three-year outcomes and thousands of events, and treat PET images as a "
     "separate, validated pipeline.", size=15)

# ================================================================ data
s = frame("Data: NACC Freeze 74 and the Cohort")
table(s, X0, TOP + 0.05, 5.0, [3.5, 1.5], [
    ["Step", "Participants"],
    ["All participants", "57,038"],
    ["First visit, in person, UDS v1–3", "55,461"],
    ["   not demented at that visit", "38,107"],
    ["   outcome known at 3 years", "19,470"],
    ["      converted to dementia", "2,818 (14.5%)"],
], size=13, row_h=0.38)
note(s, 5.8, TOP + 0.05, 3.7, 2.3, "Three groups, very different risk",
     "Normal: 12,377 (1.6% convert)\nMCI: 5,993 (41.5%)\nImpaired, not MCI: 1,100 (12.3%)", size=14)
pic(s, "split.png", X0, 3.95, W0, 2.2)
caption(s, X0, 6.35, W0, "Person-level split 70 / 15 / 15; 9 of 46 centres held out; everything learned from data is fitted on training only.")

# ================================================================ outcomes
s = frame("Outcomes and Censoring")
pic(s, "outcome.png", X0, TOP + 0.05, W0, 2.85)
note(s, X0, 4.6, 4.3, 2.4, "Primary: dementia within 3 years",
     "1 if diagnosed within 3 years; 0 if not and followed for 3 years; otherwise unknown and left out.", size=14)
note(s, 5.2, 4.6, 4.3, 2.4, "Secondary outcomes",
     "CDR-SB rise ≥ 1 within 3 years (20,355 people, 5,035 events) and the 3-year change in CDR-SB. "
     "A Cox model that keeps censored people gives the same conclusions.", size=14)

# ================================================================ features
s = frame("48 Features in Five Categories")
table(s, X0, TOP + 0.05, 4.7, [3.6, 1.1], [
    ["Category", "Features"],
    ["Demographics", "8"],
    ["Cardiovascular and medical", "25"],
    ["Cognitive tests (z-scores)", "9"],
    ["Clinical stage (CDR-SB, diagnosis)", "3"],
    ["Genetics (APOE ε4, family history)", "3"],
    ["PET (second-stage model only)", "4"],
], size=13, row_h=0.36)
note(s, 5.5, TOP + 0.05, 4.0, 2.6, "Harmonising tests across form versions",
     "Each score is compared with the expected score for the person's age, sex and education "
     "(fitted on cognitively normal people): z = (score − expected) / SD. "
     "MoCA 21 for a 72-year-old woman with 16 years of education: expected 25.8, z = −1.9.", size=13)
pic(s, "zscore.png", X0, 4.25, W0, 2.75)

# ================================================================ methodology
s = frame("Methodology: The CardioCog Pipeline")
pic(s, "pipeline.png", X0, TOP + 0.05, W0, 4.3)
note(s, X0, 5.95, W0, 1.1, "Rule for every step",
     "Only information available at the index visit is used; anything learned from data is fitted on the training "
     "set; results are reported on the test set and on centres never seen in training.", size=13)

# ================================================================ techniques
s = frame("Techniques & Tools")
box(s, X0, TOP, W0, 5.6, [
    {"text": "Models compared on the same 48 features", "h": True, "before": 0},
    {"text": "Logistic regression, random forest, **XGBoost** (main model), MLP, FT-Transformer", "lvl": 0},
    {"text": "Modality-encoder fusion network with attention, risk and progression heads", "lvl": 0},
    {"text": "Evaluation", "h": True},
    {"text": "AUROC, AUPRC, Brier score, calibration, net benefit; 1,000 bootstrap resamples", "lvl": 0},
    {"text": "Explainability and evidence", "h": True},
    {"text": "SHAP, category ablation, Grad-CAM for the PET CNN", "lvl": 0},
    {"text": "Hybrid retrieval: MiniLM embeddings + BM25, merged by rank fusion", "lvl": 0},
    {"text": "Data and platforms", "h": True},
    {"text": "NACC Freeze 74: UDS clinical visits (v1–3), PET core values and PET images", "lvl": 0},
    {"text": "Python (scikit-learn, XGBoost, PyTorch, SHAP, lifelines, SimpleITK); FastAPI + React prototype", "lvl": 0},
], size=16, space=3)

# ================================================================ fusion architecture
s = frame("Modality-Encoder Fusion Network")
pic(s, "fusion_architecture.png", X0, TOP + 0.05, W0, 3.4)
box(s, X0, 5.0, W0, 2.0, [
    "Each category has its own encoder; attention combines the five tokens into one summary.",
    "A missing category gets a learned token; whole categories are hidden with probability 0.35 in training.",
    "Loss: cross-entropy for risk + 0.5 × Huber loss for the CDR-SB change.",
], size=15, space=4)

# ================================================================ results dementia
s = frame("Results: Dementia Within Three Years")
table(s, X0, TOP + 0.05, W0, [2.3, 2.2, 1.1, 1.0, 1.4, 1.0], [
    ["Model", "Test AUROC (95% CI)", "AUPRC", "Brier", "Held-out AUROC", "Brier"],
    ["Logistic regression", "0.936 (0.924–0.949)", "0.721", "0.066", "0.945", "0.060"],
    ["Random forest", "0.932 (0.919–0.944)", "0.704", "0.070", "0.940", "0.064"],
    ["XGBoost", "0.938 (0.925–0.948)", "0.713", "0.066", "0.946", "0.059"],
    ["MLP", "0.938 (0.926–0.949)", "0.721", "0.066", "0.946", "0.059"],
    ["FT-Transformer", "0.939 (0.927–0.949)", "0.720", "0.066", "0.945", "0.061"],
], size=13, row_h=0.45, bold_cells={(3, 0), (3, 1), (3, 4)})
caption(s, X0, 4.25, W0, "Test set: 2,286 people, 332 events.  Held-out centres: 4,064 people, 545 events.")
note(s, X0, 4.85, 4.3, 2.2, "What it shows",
     "All five models discriminate well and are well calibrated, at new centres as well as on the test set.", size=15)
note(s, 5.2, 4.85, 4.3, 2.2, "What it does not show",
     "The intervals overlap almost completely: deep models give no gain over logistic regression or boosted trees.", size=15)

# ================================================================ within groups
s = frame("Is 0.94 Too Good? Results Within Groups")
table(s, X0, TOP + 0.05, 4.5, [2.5, 1, 1], [
    ["Group", "Test", "Held-out"],
    ["All not demented", "0.937", "0.946"],
    ["Cognitively normal", "0.884", "0.876"],
    ["MCI", "0.835", "0.846"],
], size=14, row_h=0.48)
caption(s, X0, 3.45, 4.5, "Converters: normal 24 / 37, MCI 289 / 496 (test / held-out).")
note(s, 5.3, TOP + 0.05, 4.2, 2.2, "Critical check",
     "Removing all current-cognition information drops AUROC from 0.937 to 0.746.", size=15)
note(s, X0, 4.3, 4.3, 2.6, "Interpretation",
     "Diagnosis alone gives AUROC 0.84, because MCI converts 25 times as often as normal cognition. "
     "Within MCI the model still ranks risk at 0.84.", size=15)
note(s, 5.3, 4.3, 4.2, 2.6, "Meaning",
     "The model estimates short-term progression given the current clinical stage. It is not leakage: inputs come "
     "from the index visit, the label from later visits.", size=15)

# ================================================================ ablation
s = frame("What Drives the Prediction?")
pic(s, "06_feature_groups.png", X0, TOP + 0.05, W0, 2.6)
table(s, 1.0, 4.3, 8.0, [3.2, 1.3, 1.9, 1.6], [
    ["Category", "Alone", "Gain added last (test)", "Gain (held-out)"],
    ["Cognitive tests", "0.896", "+0.022", "+0.019"],
    ["Clinical stage", "0.897", "+0.015", "+0.021"],
    ["Demographics", "0.671", "+0.010", "+0.009"],
    ["Genetics (APOE)", "0.645", "+0.003", "+0.001"],
    ["Cardiovascular / medical", "0.626", "+0.001", "+0.001"],
], size=13, row_h=0.42)

# ================================================================ cardiovascular
s = frame("The Cardiovascular Question (RQ2)")
pic(s, "09_hazard_ratios.png", X0, TOP + 0.05, 4.7, 3.0)
box(s, 5.4, TOP + 0.05, 4.1, 3.0, [
    "Alone: AUROC **0.626**.",
    "After cognition: **+0.001** for dementia, +0.0004 for CDR-SB decline.",
    "Cox model: hypertension HR 0.995; diabetes and stroke ≈ 1.",
], size=15, space=8)
note(s, X0, 4.9, W0, 2.0, "Interpretation",
     "Limited added value in this cohort, not evidence that vascular health is unimportant: median age 71, midlife "
     "exposure unobserved, and past damage already shows in the cognitive scores.", size=15)

# ================================================================ shap
s = frame("Explainability: Shapley Values")
pic(s, "07_shap_features.png", X0, TOP + 0.05, 5.6, 3.6)
table(s, 6.3, TOP + 0.05, 3.2, [2.1, 1.1], [
    ["Category", "SHAP share"],
    ["Clinical stage", "40.8%"],
    ["Cognitive tests", "35.0%"],
    ["Demographics", "15.5%"],
    ["Cardiovascular", "5.1%"],
    ["Genetics", "3.5%"],
], size=13, row_h=0.42)
box(s, X0, 5.35, W0, 1.7, [
    "Each prediction is split into contributions from each feature.",
    "Rankings are stable over 8 refits (Spearman 0.89–0.97).",
    "Low BMI raises the estimated risk: weight loss before diagnosis, not a causal effect.",
], size=15, space=3)

# ================================================================ decline
s = frame("Cognitive Decline as the Outcome")
pic(s, "decline_vs_dementia.png", X0, TOP + 0.05, W0, 3.3)
note(s, X0, 4.95, 4.3, 2.1, "CDR-SB rise ≥ 1",
     "20,355 people, 5,035 events. AUROC 0.858 test, 0.885 held-out; within normal cognition 0.76–0.83.", size=14)
note(s, 5.2, 4.95, 4.3, 2.1, "Predicted 3-year CDR-SB change",
     "Mean absolute error 0.69 → 0.60 (test), 0.82 → 0.72 (held-out); MCI 1.90 → 1.59; no gain for normal cognition.", size=14)

# ================================================================ fusion results
s = frame("Fusion Network: Results")
pic(s, "fusion_attention.png", X0, TOP + 0.05, W0, 3.0)
table(s, X0, 4.75, 4.2, [1.6, 1.3, 1.3], [
    ["AUROC", "Fusion net", "XGBoost"],
    ["Test", "0.934", "0.937"],
    ["Held-out", "0.943", "0.946"],
    ["MCI (test)", "0.831", "0.835"],
], size=13, row_h=0.42)
note(s, 5.1, 4.75, 4.4, 2.3, "Finding",
     "As good as XGBoost but no better. Attention is 18–22% for every category (uniform = 20%), so it does not "
     "explain predictions; SHAP and ablation are used instead.", size=14)

# ================================================================ visits and PET
s = frame("Earlier Visits and PET Measures")
pic(s, "06_longitudinal.png", X0, TOP + 0.05, 5.8, 2.0)
note(s, 6.5, TOP + 0.05, 3.0, 2.2, "Visit history (RQ5)",
     "Using up to all earlier visits adds +0.000 to +0.005 AUROC.", size=14)
pic(s, "08_conversion_by_amyloid.png", X0, 4.0, 5.8, 2.4)
note(s, 6.5, 4.0, 3.0, 3.0, "Amyloid and tau PET (RQ3)",
     "In MCI, 34% of amyloid-positive vs 7% of negative people convert within 2 years, but adding PET to the "
     "model gives no demonstrable gain (243 people, 33 events).", size=14)

# ================================================================ PET pipeline
s = frame("PET Image Pipeline: DICOM to SUVR")
pic(s, "11_group_average_images.png", X0, TOP + 0.05, 4.4, 3.4)
caption(s, X0, 4.9, 4.4, "Group averages only (318 amyloid, 232 tau, 49 FDG scans).")
box(s, 5.2, TOP + 0.05, 4.3, 3.6, [
    "**Convert:** 599 dynamic DICOM series → one static volume.",
    "**Register:** to MNI152, then to a study PET template.",
    "**Quantify:** SUVR = mean cortex ÷ mean cerebellum.",
], size=15, space=10)
note(s, X0, 5.5, W0, 1.5, "Validation against the NACC PET core",
     "Pearson r = 0.977 over 318 scans (0.96–0.99 per tracer); mean difference −0.043.", size=15)

# ================================================================ CNN
s = frame("3D CNN for Amyloid Status")
pic(s, "11_gradcam_mean.png", X0, TOP + 0.05, 4.4, 3.3)
table(s, 5.2, TOP + 0.05, 4.3, [3.0, 1.3], [
    ["Method (292 scans, 84 positive)", "AUROC"],
    ["Cortical SUVR (one number)", "0.982"],
    ["Regional logistic model", "0.971"],
    ["3D CNN", "0.960"],
], size=13, row_h=0.45)
note(s, 5.2, 3.45, 4.3, 1.5, "Interpretation",
     "The image model works, but the label itself comes from cortical SUVR, so one number is hard to beat.", size=14)
note(s, X0, 5.15, W0, 1.8, "Limitation",
     "Grad-CAM maps were diffuse, and only 5 imaged people converted, so images were not used to predict dementia.", size=14)

# ================================================================ validation
s = frame("Validation: Calibration, Usefulness, Missing Data")
pic(s, "09_decision_curve.png", X0, TOP + 0.05, W0, 2.85)
box(s, X0, 4.55, 4.2, 2.5, [
    "Predicted 14.5% = observed 14.5% (test).",
    "At 20% risk: sensitivity 83%, specificity 88%.",
    "Useful (net benefit) from 2% to 60% risk.",
], size=14, space=5)
table(s, 4.9, 4.55, 4.6, [2.0, 0.95, 0.95, 1.0], [
    ["Brier score", "Standard", "Dropout", "Specialist"],
    ["Everything", "0.066", "0.067", "0.066"],
    ["No clinical stage", "0.185", "0.076", "0.074"],
    ["No cognition at all", "0.441", "0.115", "0.111"],
    ["Demographics only", "0.444", "0.120", "0.118"],
], size=12, row_h=0.38)
caption(s, 4.9, 6.55, 4.6, "Modality dropout keeps one model close to a specialist model.")

# ================================================================ prototype
s = frame("Research Prototype for Clinicians")
pic(s, "app_estimate.png", 5.6, TOP + 0.05, 3.9, 5.6)
box(s, X0, TOP + 0.05, 4.8, 5.6, [
    {"text": "Example: 72-year-old woman, 16 years of education, MCI, CDR-SB 2, MoCA 21.", "plain": True, "bold": True, "gap": 10},
    "FastAPI + React; the form is generated from the 48 trained features.",
    "Risk shown next to the observed rate of similar people, with SHAP by category.",
    "Plausibility checks flag contradictory entries.",
    "27 of 28 behaviour checks pass.",
], size=15, space=10)

# ================================================================ report
s = frame("Plain-Language Report")
pic(s, "app_report.png", 5.4, TOP + 0.05, 4.1, 5.6)
box(s, X0, TOP + 0.05, 4.6, 5.6, [
    "**Summary:** risk next to the observed rate of similar people (39% vs 40%).",
    "**Main reasons:** largest SHAP contributions in plain words.",
    "**Reliability:** missing categories and warnings.",
    "**Published evidence:** quoted passages with their source.",
    "**Limits:** a fixed statement that it is not a diagnosis.",
    {"text": "Evidence is quoted, never generated, so the report cannot state anything the model or the literature did not provide.",
     "plain": True, "color": GREY},
], size=15, space=9)

# ================================================================ evidence assistant
s = frame("Evidence Assistant (RAG)")
pic(s, "app_ask.png", X0, TOP + 0.05, W0, 2.6)
box(s, X0, 4.35, 4.3, 2.6, [
    "MiniLM embeddings + BM25 keyword search, merged by rank fusion.",
    "If nothing is similar enough, it says the library has no evidence.",
], size=15, space=8)
table(s, 5.0, 4.35, 4.5, [3.2, 1.3], [
    ["Test (11 papers, 447 passages)", "Result"],
    ["Correct source first", "32 / 32"],
    ["Personal-advice questions flagged", "8 / 8"],
    ["Out-of-scope questions declined", "6 / 6"],
], size=13, row_h=0.45)

# ================================================================ verification
s = frame("Independent Verification")
table(s, X0, TOP + 0.05, W0, [5.2, 3.8], [
    ["Check", "Result"],
    ["No person or held-out centre in more than one partition", "0 shared"],
    ["Features unchanged when all later visits are deleted", "Identical (1,500 people × 67 features)"],
    ["No PET value newer than the visit that uses it", "0 of 1,481 amyloid, 713 tau"],
    ["Labels recomputed independently", "0 disagreements in 11,403 visits"],
    ["Shuffled labels", "AUROC 0.508 (chance)"],
    ["Retrain with a new seed", "0.937 / 0.946 (reported 0.938 / 0.946)"],
], size=13, row_h=0.45)
note(s, X0, 4.85, W0, 2.1, "Errors the checks found, and we fixed",
     "Education and race filled from a later visit (0.1% of people); PET scans matched to visits weeks before the scan. "
     "After fixing, every model was retrained; results moved only in the third decimal. 17 of 17 checks pass.", size=15)

# ================================================================ synthesis
s = frame("What Each Research Question Found")
table(s, X0, TOP + 0.05, W0, [2.9, 6.1], [
    ["Research question", "Answer"],
    ["RQ1 Prediction accuracy", "Dementia 0.94 (0.84 in MCI); decline 0.86 (0.75 in MCI)"],
    ["RQ2 Cardiovascular history", "+0.001 AUROC once cognition is known"],
    ["RQ3 Amyloid and tau PET", "Strong association, no demonstrable gain yet"],
    ["RQ4 Deep models / fusion net", "No gain over logistic regression or XGBoost"],
    ["RQ5 Earlier visits", "+0.000 to +0.005 AUROC"],
    ["RQ6 Missing categories", "Modality dropout keeps calibration close to specialist models"],
    ["RQ7 Explanation and evidence", "SHAP, plausibility checks, report with quoted evidence"],
], size=14, row_h=0.5, bold_col0=True)
note(s, X0, 5.7, W0, 1.4, "Final takeaway",
     "Current cognitive and functional status carries most of the three-year signal; vascular and genetic history "
     "add little once it is known.", size=15)

# ================================================================ limitations
s = frame("Limitations & Future Work")
box(s, X0, TOP, W0, 5.6, [
    "Held-out centres are still within NACC; validation on another cohort is the next step.",
    "NACC is a research cohort, mostly White and highly educated; about half of eligible people lacked a 3-year outcome.",
    "Few cognitively normal converters; normal cognition needs a 5- to 10-year model.",
    "Richer vascular data (repeated blood pressure, treatment, imaging) may capture what yes/no history misses.",
    "PET registered without MRI and tau not quantified; longer follow-up will allow a proper PET test.",
    "Evidence library of 11 papers; a larger curated library would help.",
], size=17, space=12)

# ================================================================ relevance
s = frame("Industry / Research Relevance")
box(s, X0, TOP, W0, 5.6, [
    "**Clinical screening:** flags elevated short-term risk for earlier specialist follow-up, with reasons and evidence, without making a diagnosis.",
    "**Clinical trials:** within MCI the model ranks risk at AUROC 0.84, which helps select people likely to progress.",
    "**Research:** a reusable method for measuring what each data source adds and for checking a pipeline against information leakage.",
], size=18, space=18)

# ================================================================ team roles
s = frame("Team Roles")
table(s, X0, TOP + 0.1, W0, [3.2, 5.8], [
    ["Members", "Responsibility"],
    ["Bhargav, Hansika", "Cardiovascular and demographic features, baseline models"],
    ["Akhil, Bhargav", "Visit-history models, CDR-SB progression regression"],
    ["Mahidhar, Hansika, Bhargav", "PET image pipeline and 3D CNN, modality-encoder fusion network"],
    ["Akhil, Hansika, Mahidhar", "RAG evidence assistant, evaluation framework, ablations"],
], size=15, row_h=0.75, bold_col0=True)

# ================================================================ references
s = frame("References")
box(s, X0, TOP, W0, 5.7, [
    "[1] G. Livingston et al., \"Dementia prevention, intervention, and care: 2020 report of the Lancet Commission,\" The Lancet, vol. 396, pp. 413–446, 2020.",
    "[2] S. A. Alaka et al., \"Enhancing the validity of CAIDE dementia risk scores with resting heart rate and machine learning: an analysis from the National Alzheimer's Coordinating Center,\" Alzheimer's & Dementia, vol. 21, e70442, 2025.",
    "[3] A. Stern and M. Linial, \"Integrative machine learning approach to risk prediction for dementia and Alzheimer's disease,\" GeroScience, vol. 48, pp. 3007–3023, 2026.",
    "[4] D. Aschwanden et al., \"Predicting cognitive impairment and dementia: a machine learning approach,\" J. Alzheimer's Disease, vol. 75, pp. 717–728, 2020.",
    "[5] F. Yi et al., \"XGBoost-SHAP-based interpretable diagnostic framework for Alzheimer's disease,\" BMC Med. Inform. Decis. Mak., vol. 23, 137, 2023.",
    "[6] Y. Gorishniy et al., \"Revisiting deep learning models for tabular data,\" NeurIPS, vol. 34, 2021.",
    "[7] P. Lewis et al., \"Retrieval-augmented generation for knowledge-intensive NLP tasks,\" NeurIPS, vol. 33, 2020.",
    "[8] T. Chen and C. Guestrin, \"XGBoost: a scalable tree boosting system,\" Proc. ACM SIGKDD, pp. 785–794, 2016.",
    "[9] S. M. Lundberg and S.-I. Lee, \"A unified approach to interpreting model predictions,\" NeurIPS, vol. 30, 2017.",
    "[10] W. E. Klunk et al., \"The Centiloid Project: standardizing quantitative amyloid plaque estimation by PET,\" Alzheimer's & Dementia, vol. 11, pp. 1–15, 2015.",
], size=12, space=5, bullets=False)

# ================================================================ thank you
s = clone(1)
retitle = shape(s, "Title 1")
place(retitle, 0.6, 2.4, 8.8, 1.3)
fill(retitle.text_frame, [{"text": "Thank You", "plain": True}], size=44, bullets=False, align=PP_ALIGN.CENTER)
sub = shape(s, "Subtitle 2")
place(sub, 0.5, 3.9, 9.0, 2.8)
fill(sub.text_frame, [
    {"text": "Questions are welcome.", "size": 20, "color": INK, "gap": 28},
    {"text": "Data: National Alzheimer's Coordinating Center (NACC), Freeze 74, used under the NACC data use agreement. All results are aggregates."},
], size=13, color=GREY, bullets=False, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- drop the template's own slides
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst)[:N_REF]:
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)

prs.save(str(OUT))
print(f"wrote {OUT.name}: {len(prs.slides)} slides")
